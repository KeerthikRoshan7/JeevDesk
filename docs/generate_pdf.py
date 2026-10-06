"""
JEEVDESK Project Document & PDF Generator
Generates a comprehensive, publication-grade document and converts it to PDF using Word COM.
"""
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import win32com.client

# Neelam Blue Theme Colors
COLOR_NEELAM_NAVY = RGBColor(0, 31, 84)     # #001F54
COLOR_NEELAM_BLUE = RGBColor(3, 64, 120)    # #034078
COLOR_NEELAM_CYAN = RGBColor(0, 119, 182)   # #0077B6
COLOR_SAFFRON_GOLD = RGBColor(217, 119, 6)  # #D97706
COLOR_DARK_TEXT = RGBColor(30, 41, 59)      # #1E293B
COLOR_MUTED_TEXT = RGBColor(100, 116, 139)  # #64748B

HEX_NEELAM_NAVY = "001F54"
HEX_NEELAM_LIGHT = "F0F7FF"
HEX_BORDER_LIGHT = "CBD5E1"
HEX_CODE_BG = "F8FAFC"

def set_cell_background(cell, hex_color):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_callout(doc, title, text, border_color="001F54", bg_color="F0F7FF"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    # Left border only
    borders_xml = f'''
    <w:tcBorders {nsdecls("w")}>
        <w:top w:val="none"/>
        <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>
        <w:bottom w:val="none"/>
        <w:right w:val="none"/>
    </w:tcBorders>
    '''
    cell._tc.get_or_add_tcPr().append(parse_xml(borders_xml))
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"✦ {title}\n")
    run_t.font.name = "Calibri"
    run_t.font.size = Pt(11)
    run_t.font.bold = True
    run_t.font.color.rgb = COLOR_NEELAM_NAVY
    
    run_b = p.add_run(text)
    run_b.font.name = "Calibri"
    run_b.font.size = Pt(10)
    run_b.font.color.rgb = COLOR_DARK_TEXT
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_code_block(doc, code_str, caption=""):
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.paragraph_format.space_before = Pt(6)
        p_cap.paragraph_format.space_after = Pt(2)
        r_cap = p_cap.add_run(f"Listing: {caption}")
        r_cap.font.name = "Calibri"
        r_cap.font.size = Pt(9.5)
        r_cap.font.bold = True
        r_cap.font.color.rgb = COLOR_NEELAM_BLUE

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, HEX_CODE_BG)
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)

    borders_xml = f'''
    <w:tcBorders {nsdecls("w")}>
        <w:top w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER_LIGHT}"/>
        <w:left w:val="single" w:sz="18" w:space="0" w:color="{HEX_NEELAM_NAVY}"/>
        <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER_LIGHT}"/>
        <w:right w:val="single" w:sz="6" w:space="0" w:color="{HEX_BORDER_LIGHT}"/>
    </w:tcBorders>
    '''
    cell._tc.get_or_add_tcPr().append(parse_xml(borders_xml))

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(code_str)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(15, 23, 42)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def generate_jeevdesk_documentation():
    doc = docx.Document()

    # Page Margins: 0.75 in (54 pt)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # ------------------ COVER PAGE ------------------
    p_pre = doc.add_paragraph()
    p_pre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_pre.paragraph_format.space_before = Pt(20)
    p_pre.paragraph_format.space_after = Pt(4)
    r_gov = p_pre.add_run("GOVERNMENT OF INDIA • MINISTRY OF SOCIAL JUSTICE AND EMPOWERMENT (MoSJE)")
    r_gov.font.name = "Calibri"
    r_gov.font.size = Pt(10)
    r_gov.font.bold = True
    r_gov.font.color.rgb = COLOR_SAFFRON_GOLD

    p_subgov = doc.add_paragraph()
    p_subgov.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_subgov.paragraph_format.space_after = Pt(16)
    r_sub = p_subgov.add_run("DEPARTMENT OF SOCIAL JUSTICE AND EMPOWERMENT | PM-AJAY GRANT-IN-AID")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(9)
    r_sub.font.color.rgb = COLOR_MUTED_TEXT

    # Logo Emblem
    logo_path = r"d:\PROJECTS\SIH2026\jeevdesk\frontend\assets\jeevdesk_emblem_1x1.jpg"
    if os.path.exists(logo_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_after = Pt(14)
        doc.add_picture(logo_path, width=Inches(1.8))
        # Center image
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run("JEEVDESK (ஜீவ்டெஸ்க்)")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(26)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_NEELAM_NAVY

    p_subtitle = doc.add_paragraph()
    p_subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_subtitle.paragraph_format.space_after = Pt(20)
    r_subt = p_subtitle.add_run("AI-Driven Multilingual Voice Assistant for Livelihood Mapping & NSQF-Aligned Skilling Recommendations")
    r_subt.font.name = "Calibri"
    r_subt.font.size = Pt(12)
    r_subt.font.bold = True
    r_subt.font.color.rgb = COLOR_NEELAM_BLUE

    # Meta Table
    tbl_meta = doc.add_table(rows=5, cols=2)
    tbl_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_rows = [
        ("Program Component", "Pradhan Mantri Anusuchit Jaati Abhyuday Yojana (PM-AJAY) Grant-in-Aid"),
        ("Target Beneficiaries", "Scheduled Caste (SC) Individuals, Self-Help Groups (SHGs) & Rural Youth"),
        ("Framework Alignment", "National Skills Qualification Framework (NSQF Levels 3 - 4)"),
        ("Primary Modalities", "Multilingual Voice-First (Tamil, Hindi, English) / Assisted Facilitator Portal"),
        ("Release & Version", "Version 1.0.0-Production (Neelam Architecture)")
    ]
    for idx, (label, val) in enumerate(meta_rows):
        c0 = tbl_meta.cell(idx, 0)
        c1 = tbl_meta.cell(idx, 1)
        c0.width = Inches(2.3)
        c1.width = Inches(4.7)
        set_cell_background(c0, HEX_NEELAM_LIGHT)
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)

        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        r0.font.name = "Calibri"
        r0.font.size = Pt(9.5)
        r0.font.bold = True
        r0.font.color.rgb = COLOR_NEELAM_NAVY

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(val)
        r1.font.name = "Calibri"
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = COLOR_DARK_TEXT

    doc.add_page_break()

    # ------------------ SECTION 1: WHAT IS JEEVDESK ------------------
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.color.rgb = COLOR_NEELAM_NAVY
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = COLOR_NEELAM_BLUE
        return p

    def add_body(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(10)
        r.font.color.rgb = COLOR_DARK_TEXT
        return p

    add_h1("1. Executive Summary & What is JEEVDESK")
    add_body(
        "JEEVDESK (ஜீவ்டெஸ்க்) is an empathetic, multilingual, voice-first artificial intelligence assistant engineered "
        "specifically for Scheduled Caste (SC) beneficiaries supported under the Grant-in-Aid (GIA) component of the "
        "Pradhan Mantri Anusuchit Jaati Abhyuday Yojana (PM-AJAY). Operating across rural districts in regional languages "
        "and local dialects (such as Tamil, Hindi, and Indian English), the system conducts structured, conversational voice "
        "interviews to construct an individual livelihood profile, identify existing vocational proficiencies, and recommend "
        "locally viable, NSQF-aligned qualifications and micro-enterprise grants."
    )

    add_callout(
        doc,
        "The Product Promise",
        "To help each beneficiary transition seamlessly from 'I need sustainable livelihood work' to a verified, "
        "locally reachable, and financially supported skilling or enterprise opportunity—completely eliminating the barrier "
        "of literacy-intensive paperwork and complex digital application portals.",
        border_color="001F54"
    )

    add_h2("1.1 Purpose & Policy Context")
    add_body(
        "Under PM-AJAY Grant-in-Aid guidelines issued by the Department of Social Justice and Empowerment (MoSJE), funds are "
        "earmarked to execute comprehensive livelihood projects that combine skill development, income-generating infrastructure, "
        "and market linkages. Historically, a significant delivery gap has existed: beneficiaries with traditional family skills "
        "(such as agriculture, tailoring, or metalcraft) were categorized as 'unskilled' on conventional portals, resulting in "
        "high course dropout rates and poor post-training placement. JEEVDESK bridges this exact gap by utilizing explainable AI "
        "to map experiential background directly to certified NSQF job roles and local FPO/enterprise vacancies."
    )

    add_h2("1.2 Target Beneficiaries & Primary Personas")
    add_body(
        "1. Rural Youth & Low-Literacy Beneficiaries: Individuals operating basic or shared smartphones who communicate exclusively in their native tongue and require intuitive voice interaction.\n"
        "2. Traditional & Informal Artisans: Beneficiaries engaged in farm labor, weaving, or repair work who need formal recognition of prior learning (RPL) and upward progression into modern vocations.\n"
        "3. Wage-Seeking Job Seekers: Individuals seeking stable monthly wages, dormitory housing, and formal employee benefits (PF/ESI) in industrial clusters.\n"
        "4. Women SHG Micro-Entrepreneurs: Groups desiring collective agro-processing grants, shed subsidies, and buy-back market linkages.\n"
        "5. Field Facilitators & Case Workers: Ground coordinators who assist in verifying profiles, resolving speech recognition ambiguities, and escorting referrals to training hubs."
    )

    # ------------------ SECTION 2: FUNCTIONALITIES & MODULES ------------------
    add_h1("2. Core Functionalities & System Modules")
    add_body(
        "JEEVDESK is structured into five cohesive architectural modules operating cooperatively from client turn to government reporting:"
    )

    tbl_modules = doc.add_table(rows=6, cols=3)
    tbl_modules.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Module Name", "Primary Responsibilities", "PRD Traceability"]
    for i, h in enumerate(headers):
        c = tbl_modules.cell(0, i)
        set_cell_background(c, HEX_NEELAM_NAVY)
        set_cell_margins(c, 80, 80, 100, 100)
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    module_data = [
        ("Voice Dialogue Manager & Speech Pipeline", "Turn-by-turn conversational flow, dialect handling, universal repairs ('repeat', 'officer help'), and WebSpeech synthesis.", "FR-01 to FR-08"),
        ("Beneficiary Profiler & Extraction Engine", "Structured field capture (education, occupation, practical skills, travel radius, mobility constraints) with audit logs.", "FR-09 to FR-14"),
        ("NSQF Multi-Criteria Recommender", "5-factor explainable scoring engine calculating percentage fit, pruning violations, and generating bilingual justifications.", "FR-15 to FR-23"),
        ("Case Management & Referral Tracker", "Lifecycle state machine tracking referrals from intake through enrollment, training completion, and placement.", "FR-24 to FR-28"),
        ("M&E Analytics & Compliance Portal", "Real-time funnel conversion metrics, district geographic spread, and data minimization audit trails.", "FR-29 to FR-30")
    ]
    for row_idx, (m_name, m_resp, m_prd) in enumerate(module_data, start=1):
        c0 = tbl_modules.cell(row_idx, 0)
        c1 = tbl_modules.cell(row_idx, 1)
        c2 = tbl_modules.cell(row_idx, 2)
        c0.width = Inches(2.0)
        c1.width = Inches(3.8)
        c2.width = Inches(1.2)
        bg = HEX_NEELAM_LIGHT if row_idx % 2 == 1 else "FFFFFF"
        for c in (c0, c1, c2):
            set_cell_background(c, bg)
            set_cell_margins(c, 60, 60, 80, 80)
        c0.paragraphs[0].add_run(m_name).font.bold = True
        c0.paragraphs[0].runs[0].font.size = Pt(9)
        c1.paragraphs[0].add_run(m_resp).font.size = Pt(9)
        c2.paragraphs[0].add_run(m_prd).font.size = Pt(9)

    doc.add_page_break()

    # ------------------ SECTION 3: TECH STACK & SPECIFICATIONS ------------------
    add_h1("3. Technical Stack & Deployment Specifications")
    add_body(
        "JEEVDESK is engineered using a robust, cloud-native stack that optimizes for low latency, zero licensing costs, "
        "and seamless portability between local field servers and cloud infrastructure."
    )

    add_body(
        "• Core Application Backend: Python 3.13+ with FastAPI for high-throughput asynchronous REST APIs and ASGI compliance.\n"
        "• Domain Modeling & Validation: Pydantic v2 domain schemas enforcing field types, constraints, and audit metadata.\n"
        "• Frontend Architecture: Responsive HTML5, Tailwind CSS with the Neelam Sapphire Theme (#001F54 / #034078), Lucide Vector Icons.\n"
        "• Speech Processing Layer: Web Speech API & Coqui/WebRTC fallback for client-side low-latency speech synthesis and recognition.\n"
        "• Data Persistence & Repositories: Local JSON Datastores with TimescaleDB / PostgreSQL readiness for production scaling.\n"
        "• Serverless Cloud Hosting: Vercel Python Serverless Runtime (@vercel/python) via vercel.json and api/index.py entrypoints."
    )

    add_h2("3.1 Recommendation Scoring Mathematical Formulation")
    add_body(
        "The recommendation engine computes a composite Total Fit Score S(b, q) for each qualification q relative to beneficiary profile b:\n\n"
        "S(b, q) = w_skill · f_skill(b, q) + w_pref · f_pref(b, q) + w_geo · f_geo(b, q) + w_edu · f_edu(b, q) + w_opp · f_opp(b, q)\n\n"
        "Where the weights are configured according to MoSJE field priorities:\n"
        "• Skill & Aspiration Affinity (w_skill = 35%): Evaluates prior tools, family occupations, and direct practical experience.\n"
        "• Livelihood Pathway Fit (w_pref = 20%): Evaluates alignment with Wage Employment vs. Self-Employment / SHG Enterprise.\n"
        "• Geographic Feasibility (w_geo = 20%): Evaluates travel distance relative to beneficiary mobility limits (15-20 km).\n"
        "• Prerequisite Compatibility (w_edu = 15%): Evaluates educational grade against NSQF entry criteria, granting credit for practical experience.\n"
        "• Market Opportunity Linkage (w_opp = 10%): Elevates qualifications tied to active, verified vacancies with FPOs or enterprises."
    )

    # ------------------ SECTION 4: DIRECTORY STRUCTURE ------------------
    add_h1("4. Project Directory Structure")
    add_body("The complete repository layout in d:\\PROJECTS\\SIH2026\\jeevdesk is organized as follows:")

    tree_str = (
        "d:\\PROJECTS\\SIH2026\\jeevdesk\\\n"
        "├── api\\\n"
        "│   └── index.py                      # Vercel Serverless ASGI Application Entrypoint\n"
        "├── backend\\\n"
        "│   └── app\\\n"
        "│       ├── main.py                   # FastAPI Application, Endpoints & CORS Middleware\n"
        "│       ├── models\\\n"
        "│       │   └── schemas.py            # Pydantic Schemas (Profile, Catalog, Rec, Case)\n"
        "│       └── services\\\n"
        "│           ├── recommender.py        # Multi-criteria NSQF Matching & XAI Scoring Engine\n"
        "│           └── voice_pipeline.py     # Multilingual Dialogue State Machine & Extraction\n"
        "├── data\\\n"
        "│   ├── qualifications_catalog.json   # 8 Approved NSQF Qualifications with PM-AJAY GIA\n"
        "│   └── opportunities_catalog.json    # 6 Verified District Opportunities, FPOs & Grants\n"
        "├── frontend\\\n"
        "│   ├── assets\\\n"
        "│   │   ├── jeevdesk_emblem_1x1.jpg   # Clean 1:1 Pixel-Perfect Official Emblem\n"
        "│   │   └── jeevdesk_logo.jpg         # Full Brand Asset\n"
        "│   └── index.html                    # Single-Page Neelam Themed Voice & Staff Console\n"
        "├── .gitignore                        # Git Exclusions for Python, Logs & Large Media\n"
        "├── README.md                         # Architecture Documentation & Deployment Guide\n"
        "├── requirements.txt                  # Python Production Dependencies\n"
        "├── run_server.py                     # Standalone Development Server Launcher\n"
        "└── vercel.json                       # Zero-Config Serverless Routing Manifest"
    )
    add_code_block(doc, tree_str, "Complete JEEVDESK Codebase Layout")

    doc.add_page_break()

    # ------------------ SECTION 5: CODE BLOCKS & EXPLANATIONS ------------------
    add_h1("5. Code Implementation & Annotated Technical Logic")
    add_body(
        "Below are the primary core modules demonstrating the production logic of JEEVDESK with in-depth annotations."
    )

    add_h2("5.1 Recommendation Engine (recommender.py)")
    add_body(
        "The recommendation engine implements Section 8.3 of the PRD. It processes beneficiary attributes, computes "
        "feature contributions, verifies constraints, and produces deterministic justifications in both English and Tamil."
    )

    code_rec = '''def match(self, profile: BeneficiaryProfile) -> RecommendationResponse:
    scored_items: List[RecommendationItem] = []

    # Map verified opportunities within beneficiary's home district
    opp_map: Dict[str, OpportunityRecord] = {
        opp.matching_qualification_id: opp 
        for opp in self.opportunities 
        if opp.district.lower() == profile.district.lower()
    }

    for qual in self.qualifications:
        factors: List[RecommendationFactor] = []

        # 1. Skill & Aspiration Fit (Weight: 35%)
        skill_score, skill_reasons = self._calculate_skill_and_occupation_affinity(profile, qual)
        factors.append(RecommendationFactor(
            factor_name="Skill & Aspiration Affinity",
            score_contribution=round(skill_score * 35, 1),
            description="; ".join(skill_reasons) if skill_reasons else "Foundational growth capability"
        ))

        # 2. Employment Preference Fit (Weight: 20%)
        pref_score, pref_desc = self._calculate_preference_fit(profile.employment_preference, qual.pathway_type)
        factors.append(RecommendationFactor(
            factor_name="Livelihood Pathway Fit",
            score_contribution=round(pref_score * 20, 1),
            description=pref_desc
        ))

        # 3. Geographic Feasibility (Weight: 20%)
        geo_score, geo_desc = self._calculate_geographic_fit(profile.district, qual.district, profile.max_travel_distance_km)
        factors.append(RecommendationFactor(
            factor_name="Geographic Feasibility",
            score_contribution=round(geo_score * 20, 1),
            description=geo_desc
        ))

        # 4. Education & Prerequisite Fit (Weight: 15%)
        edu_score, edu_desc = self._calculate_education_fit(profile.education_level, qual.prerequisites)
        factors.append(RecommendationFactor(
            factor_name="Prerequisite Compatibility",
            score_contribution=round(edu_score * 15, 1),
            description=edu_desc
        ))

        # 5. Local Market Demand Linkage (Weight: 10%)
        matched_opp = opp_map.get(qual.qualification_id)
        opp_score = 1.0 if matched_opp else 0.5
        factors.append(RecommendationFactor(
            factor_name="Market Linkage & Placement Access",
            score_contribution=round(opp_score * 10, 1),
            description=f"Verified local vacancy: {matched_opp.title}" if matched_opp else f"Cluster demand: {qual.demand_status}"
        ))

        total_fit = sum(f.score_contribution for f in factors)
        fit_pct = int(min(max(total_fit, 10), 98))

        # Bilingual plain-language explanations
        plain_en = f"Recommended because your background in {profile.current_occupation} matches this NSQF Level {qual.nsqf_level} course."
        plain_ta = f"உங்கள் அனுபவம் ({profile.current_occupation}) மற்றும் வேலை விருப்பம் ({profile.employment_preference}) இதற்கு மிகச் சிறப்பாக பொருந்துகிறது."

        scored_items.append(RecommendationItem(
            rank=0, qualification=qual, matching_opportunity=matched_opp,
            total_fit_score=total_fit, fit_percentage=fit_pct, factors=factors,
            plain_language_explanation=plain_en, regional_language_explanation=plain_ta,
            prerequisites_met=(edu_score >= 0.7), distance_feasible=(geo_score >= 0.7),
            actionable_next_step=f"Contact {qual.center_name} to enroll in the next PM-AJAY cohort."
        ))

    # Present at most 3 top pathways (PRD Section 7.3 Rule FR-19)
    scored_items.sort(key=lambda x: x.total_fit_score, reverse=True)
    top_3 = scored_items[:3]
    for i, item in enumerate(top_3):
        item.rank = i + 1

    return RecommendationResponse(
        beneficiary_id=profile.internal_id, timestamp=profile.created_at,
        recommendations=top_3, confidence_level="High", requires_human_escalation=False
    )'''
    add_code_block(doc, code_rec, "Multi-Factor NSQF Recommendation Matching Engine")

    add_h2("5.2 Multilingual Voice Dialogue Manager (voice_pipeline.py)")
    add_body(
        "The dialogue manager orchestrates conversational state transitions, handles beneficiary repair requests "
        "('repeat', 'connect to officer'), and populates the structured draft profile turn-by-turn."
    )

    code_voice = '''def process_utterance(self, session_id: str, beneficiary_utterance: str) -> Dict[str, Any]:
    session = self.sessions.get(session_id) or self.start_session()
    text = beneficiary_utterance.strip().lower()

    # Conversational Repair: Repetition Handling
    if any(w in text for w in ["repeat", "மறுபடியும்", "மீண்டும்", "दोबारा"]):
        config = CONVERSATION_FLOW_CONFIG.get(session.stage)
        lang_key = f"prompt_{session.language}"
        reply = f"(மீண்டும் சொல்கிறேன்) {config[lang_key]}" if session.language == "ta" else f"(Repeating) {config[lang_key]}"
        return {"session": session, "system_reply": reply, "stage": session.stage}

    # Conversational Repair: Human Escalation Trigger
    if any(w in text for w in ["help", "human", "அதிகாரி", "உதவி", "मदद", "officer"]):
        session.escalated_to_staff = True
        session.escalation_reason = "Beneficiary requested live human assistance during voice intake"
        reply = "உங்கள் அழைப்பு கள ஒருங்கிணைப்பாளர் உதவிப் பிரிவுக்கு மாற்றப்படுகிறது." if session.language == "ta" else "Routing to Field Facilitator queue."
        return {"session": session, "system_reply": reply, "stage": "escalated"}

    # Turn-by-Turn Field Extraction (District, Education, Skills, Preference)
    if session.stage == "district":
        for d in ["salem", "madurai", "tiruchirappalli", "namakkal", "tirupur", "dharmapuri"]:
            if d in text:
                session.profile_draft["district"] = d.capitalize()
                break
        next_stage = "current_work"

    elif session.stage == "work_preference":
        if any(w in text for w in ["சுயதொழில்", "சொந்த", "business", "self", "स्वरोजगार"]):
            session.profile_draft["employment_preference"] = "Self-Employment / Enterprise"
        else:
            session.profile_draft["employment_preference"] = "Wage Employment"
        next_stage = "completed"

    # Advance state machine and return next localized prompt
    session.stage = next_stage
    sys_reply = CONVERSATION_FLOW_CONFIG[next_stage][f"prompt_{session.language}"]
    return {"session": session, "system_reply": sys_reply, "stage": next_stage}'''
    add_code_block(doc, code_voice, "Conversational Voice Pipeline with Dialect Repair Handling")

    # Save DOCX
    docx_output_path = r"d:\PROJECTS\SIH2026\jeevdesk\docs\JEEVDESK_PROJECT_DOCUMENTATION.docx"
    os.makedirs(os.path.dirname(docx_output_path), exist_ok=True)
    doc.save(docx_output_path)
    print(f"DOCX created: {docx_output_path} (Size: {os.path.getsize(docx_output_path)} bytes)")

    # Convert to PDF via Word COM
    pdf_output_path = r"d:\PROJECTS\SIH2026\jeevdesk\docs\JEEVDESK_PROJECT_DOCUMENTATION.pdf"
    root_pdf_path = r"d:\PROJECTS\SIH2026\JEEVDESK_PROJECT_DOCUMENTATION.pdf"
    
    print("Converting DOCX to PDF using Word.Application COM...")
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        doc_obj = word.Documents.Open(os.path.abspath(docx_output_path))
        # 17 represents wdExportFormatPDF in Word COM
        doc_obj.ExportAsFixedFormat(
            OutputFileName=os.path.abspath(pdf_output_path),
            ExportFormat=17,
            OpenAfterExport=False,
            OptimizeFor=0, # wdExportOptimizeForPrint
            CreateBookmarks=1 # wdExportCreateHeadingBookmarks
        )
        doc_obj.ExportAsFixedFormat(
            OutputFileName=os.path.abspath(root_pdf_path),
            ExportFormat=17,
            OpenAfterExport=False,
            OptimizeFor=0,
            CreateBookmarks=1
        )
        doc_obj.Close()
        print(f"PDF generated successfully at: {pdf_output_path}")
        print(f"PDF copied to project root at: {root_pdf_path}")
    finally:
        word.Quit()

if __name__ == "__main__":
    generate_jeevdesk_documentation()
